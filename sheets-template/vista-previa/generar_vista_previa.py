#!/usr/bin/env python3
"""
Genera una VISTA PREVIA en .xlsx de la plantilla Gig Finance.

No es el producto: es una maqueta navegable para subir a Google Drive y ver
cómo se ve y cómo se siente la plantilla. El producto real se genera con el
proyecto de Apps Script de ../apps-script (ver ../README.md), porque hay
cosas que un .xlsx sencillamente no puede llevar: el script (marcar un pago y
que se registre el gasto), las casillas de verificación, y las fórmulas
propias de Google Sheets (QUERY, FILTER, SORT, ARRAYFORMULA, LET).

Lo que SÍ queda vivo aquí: los totales, el gasto por categoría, las barras de
presupuesto, el prorrateo de gastos recurrentes y el progreso — todo con
fórmulas reales que se recalculan al cambiar los datos.
"""

import calendar
import datetime as dt
import pathlib
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------- #
# Paleta (la misma del producto)
# --------------------------------------------------------------------- #
C_FONDO = "EEF3F9"
C_TARJETA = "FFFFFF"
C_TEXTO = "0F2547"
C_TENUE = "6B7F9E"
C_BORDE = "D7E0EC"
C_AZUL = "2563EB"
C_AZUL_SUAVE = "DBE6FB"
C_TEAL = "0D9488"
C_VERDE = "16A34A"
C_VERDE_SUAVE = "DCF5E3"
C_AMBAR = "F59E0B"
C_ROJO = "DC2626"
C_ROJO_SUAVE = "FBDEDE"
C_INPUT = "FFF8E1"

FUENTE = "Arial"
DINERO = '"$"#,##0.00'
FECHA = "dd/mm/yyyy"

fina = Side(style="thin", color=C_BORDE)
CAJA = Border(left=fina, right=fina, top=fina, bottom=fina)


def relleno(hexa):
    return PatternFill("solid", fgColor=hexa)


# openpyxl escribe las fórmulas sin valor en caché, así que el visor de Google
# Drive las muestra en blanco hasta convertir el archivo a Sheets. Vamos
# apuntando aquí el resultado de cada fórmula (hoja → celda → valor) para
# inyectarlo al final sin quitar la fórmula.
CACHE = {}


def recordar(hoja_idx, celda, valor):
    CACHE.setdefault(hoja_idx, {})[celda] = valor


def pon(ws, celda, valor=None, *, negrita=False, tam=10, color=C_TEXTO,
        fondo=None, fmt=None, alineacion=None, ajustar=False, borde=False,
        cursiva=False):
    c = ws[celda]
    if valor is not None:
        c.value = valor
    c.font = Font(name=FUENTE, bold=negrita, size=tam, color=color, italic=cursiva)
    if fondo:
        c.fill = relleno(fondo)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(
        horizontal=alineacion, vertical="center", wrap_text=ajustar
    )
    if borde:
        c.border = CAJA
    return c


def pintar(ws, rango, *, fondo=None, borde=False, fmt=None, negrita=False,
           tam=10, color=C_TEXTO, alineacion=None, ajustar=False):
    for fila in ws[rango]:
        for c in fila:
            c.font = Font(name=FUENTE, bold=negrita, size=tam, color=color)
            if fondo:
                c.fill = relleno(fondo)
            if borde:
                c.border = CAJA
            if fmt:
                c.number_format = fmt
            c.alignment = Alignment(horizontal=alineacion, vertical="center",
                                    wrap_text=ajustar)


def lienzo(ws, filas, cols):
    """Pinta el fondo base de la hoja."""
    for f in range(1, filas + 1):
        for c in range(1, cols + 1):
            cel = ws.cell(row=f, column=c)
            cel.fill = relleno(C_FONDO)
            cel.font = Font(name=FUENTE, size=10, color=C_TEXTO)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3


def banner(ws, celda_ini, celda_fin, texto):
    ws.merge_cells(f"{celda_ini}:{celda_fin}")
    pon(ws, celda_ini, texto, negrita=True, tam=10, color="7A4B00",
        fondo="FFF3CD", ajustar=True, borde=True, alineacion="center")


# --------------------------------------------------------------------- #
# Datos de ejemplo
# --------------------------------------------------------------------- #
ANIO, MES = 2026, 3
DIAS_MES = calendar.monthrange(ANIO, MES)[1]
HOY = dt.date(2026, 3, 17)
NOMBRE_MES = "Marzo 2026"

PAGOS_FIJOS = [
    ("Renta", 1200.00, 5),
    ("Seguro del carro", 140.00, 8),
    ("Luz", 85.00, 12),
    ("Internet", 60.00, 15),
    ("Celular", 45.00, 20),
]
TARJETAS = [
    ("Visa", 2400.00, 24.99, 75.00, 12),
    ("Capital One", 850.00, 29.99, 35.00, 22),
]
RECURRENTES = [
    ("Gasolina", 28.00, "Diario"),
    ("Comida en ruta", 90.00, "Semanal"),
    ("Lavado del carro", 25.00, "Quincenal"),
]
CATEGORIAS = [
    ("Gasolina", 900.00),
    ("Comida", 500.00),
    ("Mantenimiento", 200.00),
    ("Personal", 300.00),
]
PLATAFORMAS = ["Uber", "DoorDash", "Instacart", "Lyft"]

# Pagos ya marcados como pagados este mes (la Renta sigue pendiente y vencida).
PAGADOS = {"Seguro del carro", "Luz", "Visa (mínimo)"}

DIAS_TRABAJADOS = [1, 2, 4, 5, 7, 8, 9, 11, 12, 14, 15, 16]
DIAS_PLANEADOS = [18, 19, 21, 22, 23, 25, 26, 28, 29, 30]

INGRESOS = [
    (1, 145.00, "Turno del domingo", "Uber"),
    (2, 118.00, "Turno de la tarde", "DoorDash"),
    (4, 162.00, "Turno largo", "Uber"),
    (5, 96.00, "Medio turno", "Instacart"),
    (7, 188.00, "Sábado completo", "Uber"),
    (8, 155.00, "Domingo", "DoorDash"),
    (9, 104.00, "Turno corto", "Lyft"),
    (11, 131.00, "Turno de la noche", "Uber"),
    (12, 90.00, "Medio turno", "Instacart"),
    (14, 176.00, "Sábado completo", "Uber"),
    (15, 98.00, "Domingo tranquilo", "DoorDash"),
    (16, 87.00, "Turno corto", "Lyft"),
]
GASTOS_MANUALES = [
    (1, 32.00, "Gasolina", "Gasolina"), (2, 28.00, "Gasolina", "Gasolina"),
    (4, 35.00, "Gasolina", "Gasolina"), (5, 30.00, "Gasolina", "Gasolina"),
    (7, 41.00, "Gasolina", "Gasolina"), (9, 38.00, "Gasolina", "Gasolina"),
    (11, 34.00, "Gasolina", "Gasolina"), (12, 29.00, "Gasolina", "Gasolina"),
    (14, 40.00, "Gasolina", "Gasolina"), (16, 73.00, "Gasolina", "Gasolina"),
    (2, 65.00, "Comida en ruta", "Comida"), (5, 58.00, "Comida en ruta", "Comida"),
    (8, 72.00, "Comida en ruta", "Comida"), (11, 61.00, "Comida en ruta", "Comida"),
    (14, 84.00, "Comida en ruta", "Comida"), (16, 80.00, "Comida en ruta", "Comida"),
    (6, 45.00, "Cambio de aceite", "Mantenimiento"),
    (13, 185.00, "Llanta nueva", "Mantenimiento"),
    (3, 40.00, "Corte de pelo", "Personal"),
    (10, 50.00, "Salida familiar", "Personal"),
]


def prorratear(monto, frecuencia, dias):
    return {"Diario": monto * dias,
            "Semanal": monto * dias / 7,
            "Quincenal": monto * dias / 14}[frecuencia]


RECURRENTE_TOTAL = sum(prorratear(m, f, DIAS_MES) for _, m, f in RECURRENTES)

# Lista de "Pagos del mes", ordenada por día de vencimiento igual que la
# fórmula derramada del producto.
LISTA_PAGOS = (
    [(n, m, d, "Pago fijo") for n, m, d in PAGOS_FIJOS]
    + [(f"{n} (mínimo)", mi, d, "Tarjeta") for n, _, _, mi, d in TARJETAS]
)
LISTA_PAGOS.sort(key=lambda p: p[2])


def calcular_meta_diaria():
    """Mismo algoritmo que src/lib/calc.ts y que las fórmulas de la plantilla."""
    ingreso = sum(m for _, m, _, _ in INGRESOS)
    hoy = HOY

    obligaciones = []
    for nombre, monto, dia, _tipo in LISTA_PAGOS:
        if nombre in PAGADOS:
            continue
        vence = dt.date(ANIO, MES, min(dia, DIAS_MES))
        obligaciones.append((max(vence, hoy), monto, nombre))
    if RECURRENTE_TOTAL > 0:
        obligaciones.append(
            (max(dt.date(ANIO, MES, DIAS_MES), hoy), RECURRENTE_TOTAL, "Gastos recurrentes")
        )

    dias_con_ingreso = {d for d, _, _, _ in INGRESOS}
    disponibles = [
        dt.date(ANIO, MES, d) for d in sorted(set(DIAS_TRABAJADOS + DIAS_PLANEADOS))
        if dt.date(ANIO, MES, d) >= hoy and d not in dias_con_ingreso
    ]

    meta, fecha_gana, etiquetas, riesgo = 0.0, None, [], False
    for vence in sorted({o[0] for o in obligaciones}):
        acumulado = sum(m for f, m, _ in obligaciones if f <= vence)
        requerido = max(0.0, acumulado - ingreso)
        if requerido == 0:
            continue
        dias = len([d for d in disponibles if d <= vence])
        en_riesgo = dias == 0
        ritmo = requerido if en_riesgo else requerido / dias
        if ritmo > meta:
            meta, riesgo = ritmo, en_riesgo
            fecha_gana = vence
            etiquetas = [n for f, _, n in obligaciones if f == vence]
    return meta, fecha_gana, etiquetas, riesgo, len(disponibles)


META, FECHA_META, ETIQUETAS_META, EN_RIESGO, DIAS_DISPONIBLES = calcular_meta_diaria()

# --------------------------------------------------------------------- #
wb = Workbook()

# ===================================================================== #
# Hoja: Instrucciones
# ===================================================================== #
ws = wb.active
ws.title = "Instrucciones"
lienzo(ws, 90, 9)
ws.column_dimensions["B"].width = 30
for col in "CDEFGH":
    ws.column_dimensions[col].width = 15

ws.merge_cells("B2:H2")
pon(ws, "B2", "💰  Gig Finance — Presupuesto para trabajadores gig",
    negrita=True, tam=20)
ws.row_dimensions[2].height = 42

banner(ws, "B4", "H5",
       "⚠️  ESTO ES UNA VISTA PREVIA.  Sirve para ver el diseño y navegar la "
       "plantilla, no para usarla de verdad.\nLa plantilla real es un archivo "
       "nativo de Google Sheets con automatizaciones. Abajo dice cómo generarla.")
ws.merge_cells("B4:H5")
ws.row_dimensions[4].height = 30
ws.row_dimensions[5].height = 30

GUIA = [
    ("h2", "Qué SÍ funciona en esta vista previa"),
    ("p", "El diseño completo, las cuatro pestañas y los datos de ejemplo de un mes real."),
    ("p", "Los totales, el gasto por categoría, las barras de presupuesto y el prorrateo "
          "de gastos recurrentes: son fórmulas de verdad y se recalculan si cambias los datos."),
    ("h2", "Qué NO funciona aquí (y sí en la plantilla real)"),
    ("p", "Marcar un pago y que el gasto se registre solo. Necesita Apps Script, que un "
          "archivo subido a Drive no puede llevar."),
    ("p", "Las casillas de verificación reales. Aquí se ven como una ✔ escrita a mano."),
    ("p", "La meta diaria calculada al vuelo. Aquí es un número fijo de ejemplo; en la "
          "plantilla real se recalcula sola al marcar días o pagos."),
    ("p", "La lista de pagos que se arma y se ordena sola desde Configuración."),
    ("p", "El menú «Gig Finance» para crear el mes siguiente en dos clics."),
    ("h2", "Cómo generar la plantilla real (unos 5 minutos)"),
    ("num", "1.  Entra a sheets.google.com y crea una hoja de cálculo nueva y vacía."),
    ("num", "2.  Menú Extensiones ▸ Apps Script."),
    ("num", "3.  Copia ahí los archivos .gs de la carpeta apps-script del repositorio, "
            "uno por archivo, con el mismo nombre."),
    ("num", "4.  Arriba, elige la función «construirPlantilla» y dale ▶ Ejecutar. "
            "Autoriza cuando Google lo pida."),
    ("num", "5.  Vuelve a la hoja: ya está todo armado, con formato y protegido."),
    ("h2", "Las cuatro pestañas"),
    ("def", "Configuración|Lo que no cambia mes a mes: pagos fijos, tarjetas, gastos "
            "recurrentes, categorías con presupuesto y plataformas. Se llena una vez."),
    ("def", "Marzo 2026|Ejemplo de una hoja de mes: registro de movimientos, pagos del "
            "mes, calendario de trabajo y panel de resumen. En la plantilla real se llama "
            "«MES – plantilla» y se duplica cada mes."),
    ("def", "Plan de pago de deudas|Calculadora de bola de nieve / avalancha. No depende del mes."),
    ("h2", "Cómo se calcula la meta diaria"),
    ("p", "No se reparte parejo entre los días del mes. Se busca el vencimiento más "
          "exigente y se calcula el ritmo necesario para llegar a ése; cumpliendo ese "
          "ritmo, los demás vencimientos salen solos."),
    ("p", "Un pago vencido y sin marcar cuenta como que vence HOY: urgencia máxima. Y si "
          "un vencimiento no tiene ningún día de trabajo marcado antes de su fecha, la "
          "plantilla avisa «EN RIESGO» con el nombre y la fecha."),
    ("h2", "Mini-glosario"),
    ("def", "Meta diaria|Lo que necesitas generar cada día que trabajes para llegar al "
            "vencimiento más apretado que tienes por delante."),
    ("def", "Día disponible|Un día marcado en el calendario, de hoy en adelante, en el que "
            "todavía no registraste ingreso."),
    ("def", "Bola de nieve|Pagar primero la tarjeta con el SALDO MÁS CHICO. Motiva porque "
            "ves tarjetas liquidarse rápido."),
    ("def", "Avalancha|Pagar primero la tarjeta con el INTERÉS MÁS ALTO. Es la que menos "
            "intereses cuesta en total."),
    ("def", "Prorratear|Repartir un gasto frecuente para saber cuánto suma al mes. "
            "Diario × días del mes; Semanal × (días ÷ 7); Quincenal × (días ÷ 14)."),
]

f = 7
for tipo, texto in GUIA:
    if tipo == "h2":
        f += 1
        ws.merge_cells(f"B{f}:H{f}")
        pon(ws, f"B{f}", texto, negrita=True, tam=13, color=C_AZUL)
        ws.row_dimensions[f].height = 30
    elif tipo in ("p", "num"):
        ws.merge_cells(f"B{f}:H{f}")
        pon(ws, f"B{f}", texto if tipo == "num" else "▸   " + texto,
            fondo=C_TARJETA, ajustar=True, borde=True)
        ws.row_dimensions[f].height = 15 + 13 * (len(texto) // 95 + 1)
    else:
        titulo, cuerpo = texto.split("|")
        pon(ws, f"B{f}", titulo, negrita=True, fondo=C_TARJETA, ajustar=True, borde=True)
        ws.merge_cells(f"C{f}:H{f}")
        pon(ws, f"C{f}", cuerpo, fondo=C_TARJETA, ajustar=True, borde=True)
        ws.row_dimensions[f].height = 15 + 13 * (len(cuerpo) // 85 + 1)
    f += 1

f += 1
ws.merge_cells(f"B{f}:H{f}")
pon(ws, f"B{f}", "Esta plantilla es una herramienta de organización personal. "
    "No es asesoría financiera, contable ni fiscal.",
    color=C_TENUE, cursiva=True, ajustar=True)

# ===================================================================== #
# Hoja: Configuración
# ===================================================================== #
cf = wb.create_sheet("Configuración")
lienzo(cf, 95, 8)
for col, ancho in {"B": 30, "C": 17, "D": 25, "E": 17, "F": 22, "G": 8}.items():
    cf.column_dimensions[col].width = ancho

cf.merge_cells("B2:F2")
pon(cf, "B2", "⚙️  Configuración", negrita=True, tam=18)
cf.row_dimensions[2].height = 40
cf.merge_cells("B3:F3")
pon(cf, "B3", "Se llena una sola vez. Las celdas amarillas son tuyas; "
    "todo lo demás se calcula solo.", color=C_TENUE, ajustar=True)


def seccion(ws_, fila_titulo, texto, encabezados, col_ini=2):
    ws_.merge_cells(start_row=fila_titulo, start_column=2,
                    end_row=fila_titulo, end_column=6)
    pon(ws_, f"B{fila_titulo}", texto, negrita=True, tam=12, color=C_AZUL)
    ws_.row_dimensions[fila_titulo].height = 26
    for i, h in enumerate(encabezados):
        c = ws_.cell(row=fila_titulo + 1, column=col_ini + i, value=h)
        c.font = Font(name=FUENTE, bold=True, size=9, color="FFFFFF")
        c.fill = relleno(C_TEXTO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CAJA
    ws_.row_dimensions[fila_titulo + 1].height = 32


# 1 · Pagos fijos  (encabezado 5, datos 6..25)
seccion(cf, 4, "1 · Pagos fijos del mes",
        ["Nombre del pago", "Monto", "Día de vencimiento (1-31)"])
for i, (n, m, d) in enumerate(PAGOS_FIJOS):
    cf.cell(row=6 + i, column=2, value=n)
    cf.cell(row=6 + i, column=3, value=m)
    cf.cell(row=6 + i, column=4, value=d)
pintar(cf, "B6:D25", fondo=C_INPUT, borde=True)
pintar(cf, "C6:C25", fondo=C_INPUT, borde=True, fmt=DINERO)
pintar(cf, "D6:D25", fondo=C_INPUT, borde=True, alineacion="center")
pon(cf, "B26", "Total mensual de pagos fijos:", negrita=True)
pon(cf, "C26", "=SUM(C6:C25)", negrita=True, color=C_AZUL, fmt=DINERO)
recordar(2, "C26", round(sum(m for _, m, _ in PAGOS_FIJOS), 2))

# 2 · Tarjetas  (encabezado 28, datos 29..38)
seccion(cf, 27, "2 · Tarjetas de crédito",
        ["Nombre de la tarjeta", "Saldo actual", "Tasa de interés anual (APR %)",
         "Pago mínimo", "Día de vencimiento (1-31)"])
for i, (n, s, apr, mi, d) in enumerate(TARJETAS):
    cf.cell(row=29 + i, column=2, value=n)
    cf.cell(row=29 + i, column=3, value=s)
    cf.cell(row=29 + i, column=4, value=apr)
    cf.cell(row=29 + i, column=5, value=mi)
    cf.cell(row=29 + i, column=6, value=d)
pintar(cf, "B29:F38", fondo=C_INPUT, borde=True)
pintar(cf, "C29:C38", fondo=C_INPUT, borde=True, fmt=DINERO)
pintar(cf, "D29:D38", fondo=C_INPUT, borde=True, fmt='0.00"%"', alineacion="center")
pintar(cf, "E29:E38", fondo=C_INPUT, borde=True, fmt=DINERO)
pintar(cf, "F29:F38", fondo=C_INPUT, borde=True, alineacion="center")
pon(cf, "B39", "El APR se escribe como número: 24.99 significa 24.99 % anual.",
    color=C_TENUE, cursiva=True)

# 3 · Gastos recurrentes  (encabezado 42, datos 43..52)
seccion(cf, 41, "3 · Gastos recurrentes  (gasolina, comida, datos…)",
        ["Nombre", "Monto por ocurrencia", "Frecuencia"])
pon(cf, "E42", "Costo mensual aprox.", negrita=True, tam=9, color="FFFFFF",
    fondo=C_TEXTO, alineacion="center", ajustar=True, borde=True)
for i, (n, m, fr) in enumerate(RECURRENTES):
    cf.cell(row=43 + i, column=2, value=n)
    cf.cell(row=43 + i, column=3, value=m)
    cf.cell(row=43 + i, column=4, value=fr)
pintar(cf, "B43:D52", fondo=C_INPUT, borde=True)
pintar(cf, "C43:C52", fondo=C_INPUT, borde=True, fmt=DINERO)
for r in range(43, 53):
    cf[f"E{r}"] = (f'=IF($B{r}="","",$C{r}*IF($D{r}="Diario",30,'
                   f'IF($D{r}="Semanal",30/7,IF($D{r}="Quincenal",30/14,0))))')
    i_rec = r - 43
    if i_rec < len(RECURRENTES):
        _, _m, _f = RECURRENTES[i_rec]
        recordar(2, f"E{r}", round(prorratear(_m, _f, 30), 2))
    else:
        recordar(2, f"E{r}", "")
pintar(cf, "E43:E52", fondo="F5F8FC", borde=True, fmt=DINERO)
pon(cf, "B53", "Aquí se estima sobre 30 días. La hoja del mes usa los días exactos "
    "de ese mes.", color=C_TENUE, cursiva=True)

# 4 · Categorías  (encabezado 56, datos 57..71)
seccion(cf, 55, "4 · Categorías con presupuesto", ["Categoría", "Presupuesto mensual"])
for i, (n, m) in enumerate(CATEGORIAS):
    cf.cell(row=57 + i, column=2, value=n)
    cf.cell(row=57 + i, column=3, value=m)
pintar(cf, "B57:C71", fondo=C_INPUT, borde=True)
pintar(cf, "C57:C71", fondo=C_INPUT, borde=True, fmt=DINERO)

# 5 · Plataformas  (encabezado 75, datos 76..85)
seccion(cf, 74, "5 · Plataformas donde generas ingreso", ["Plataforma"])
for i, n in enumerate(PLATAFORMAS):
    cf.cell(row=76 + i, column=2, value=n)
pintar(cf, "B76:B85", fondo=C_INPUT, borde=True)
pon(cf, "B86", "Ejemplos: Uber, DoorDash, Rappi, Lyft, Instacart, Cliente directo…",
    color=C_TENUE, cursiva=True)

# ===================================================================== #
# Hoja: Marzo 2026
# ===================================================================== #
ms = wb.create_sheet(NOMBRE_MES)
FILA_MOV_1, FILA_MOV_N = 16, 120        # la real llega a 515
FILA_PAG_1, FILA_PAG_N = 16, 45
CAL_F1 = 16
PRE_F1, PRE_FN = 16, 30
lienzo(ms, 135, 48)

anchos = {"B": 13, "C": 12, "D": 13, "E": 27, "F": 17, "G": 19, "H": 2,
          "I": 2, "J": 10, "K": 27, "L": 13, "M": 7, "N": 12, "O": 2,
          "P": 8, "Q": 8, "R": 8, "S": 8, "T": 8, "U": 8, "V": 8, "W": 2,
          "X": 20, "Y": 14, "Z": 14, "AA": 7, "AB": 20}
for col, a in anchos.items():
    ms.column_dimensions[col].width = a

ms.merge_cells("B2:H2")
pon(ms, "B2", f"📅  {NOMBRE_MES} — Gig Finance", negrita=True, tam=18)
ms.row_dimensions[2].height = 40
pon(ms, "J2", "Mes:", negrita=True, alineacion="right")
pon(ms, "K2", "Marzo", negrita=True, fondo=C_INPUT, borde=True, alineacion="center")
pon(ms, "L2", "Año:", negrita=True, alineacion="right")
pon(ms, "M2", 2026, negrita=True, fondo=C_INPUT, borde=True, alineacion="center")
ms.merge_cells("P2:V2")
pon(ms, "P2", "Vista previa · datos de ejemplo", color=C_TENUE, cursiva=True,
    ajustar=True, alineacion="center")

# ---- Zona de cálculo (oculta, como en el producto) ----
ms["AE5"] = DIAS_MES
ms["AE8"] = f'=SUMIFS($D${FILA_MOV_1}:$D${FILA_MOV_N},$C${FILA_MOV_1}:$C${FILA_MOV_N},"Ingreso")'
ms["AE9"] = f'=SUMIFS($D${FILA_MOV_1}:$D${FILA_MOV_N},$C${FILA_MOV_1}:$C${FILA_MOV_N},"Gasto")'
for i in range(len(RECURRENTES)):
    r = 43 + i
    ms.cell(row=1 + i, column=34).value = (
        f"=IF(Configuración!$B{r}=\"\",0,Configuración!$C{r}*"
        f'IF(Configuración!$D{r}="Diario",$AE$5,'
        f'IF(Configuración!$D{r}="Semanal",$AE$5/7,'
        f'IF(Configuración!$D{r}="Quincenal",$AE$5/14,0))))'
    )
ms["AE10"] = f"=SUM(AH1:AH{max(1, len(RECURRENTES))})"
_ING = round(sum(m for _, m, _, _ in INGRESOS), 2)
_GAS = round(sum(m for _, m, _, _ in GASTOS_MANUALES)
             + sum(m for n, m, _, _ in LISTA_PAGOS if n in PAGADOS), 2)
_TOTAL_PAGOS = round(sum(m for _, m, _, _ in LISTA_PAGOS), 2)
_SIN_MARCAR = round(sum(m for n, m, _, _ in LISTA_PAGOS if n not in PAGADOS), 2)
_OBLIG = round(_TOTAL_PAGOS + RECURRENTE_TOTAL, 2)
_FALTA = round(max(0, _SIN_MARCAR + RECURRENTE_TOTAL - _ING), 2)
_PROGRESO = min(100, round(_ING / _OBLIG * 100)) if _OBLIG > 0 else 100
for _i, (_n, _m, _f) in enumerate(RECURRENTES):
    recordar(3, f"AH{_i+1}", round(prorratear(_m, _f, DIAS_MES), 2))
for _celda, _v in [("AE8", _ING), ("AE9", _GAS), ("AE10", round(RECURRENTE_TOTAL, 2)),
                   ("AE12", _OBLIG), ("AE19", _SIN_MARCAR), ("AE13", _FALTA),
                   ("AE14", _PROGRESO)]:
    recordar(3, _celda, _v)
ms["AE12"] = f"=SUM($L${FILA_PAG_1}:$L${FILA_PAG_N})+$AE$10"
ms["AE19"] = (f'=SUMIFS($L${FILA_PAG_1}:$L${FILA_PAG_N},'
              f'$J${FILA_PAG_1}:$J${FILA_PAG_N},"",'
              f'$K${FILA_PAG_1}:$K${FILA_PAG_N},"<>")')
ms["AE13"] = "=MAX(0,$AE$19+$AE$10-$AE$8)"
ms["AE14"] = "=IF($AE$12>0,MIN(100,ROUND($AE$8/$AE$12*100,0)),100)"
ms["AE15"] = round(META, 2)          # estático: en el producto es fórmula viva
for etiqueta, celda in [("días del mes", "AD5"), ("ingresos", "AD8"),
                        ("gastos", "AD9"), ("recurrentes", "AD10"),
                        ("total obligaciones", "AD12"), ("falta por cubrir", "AD13"),
                        ("progreso %", "AD14"), ("meta diaria (fija en la vista previa)", "AD15"),
                        ("pagos sin marcar", "AD19")]:
    ms[celda] = etiqueta

# ---- Panel ----
ms.merge_cells("B4:V4")
pon(ms, "B4", "PANEL DEL MES", negrita=True, tam=11, color=C_AZUL)
ms.row_dimensions[4].height = 26


def tarjeta_kpi(fila, col, etiqueta, formula, fmt=DINERO, color=C_TEXTO, tam=15,
                valor=None):
    c1, c2 = get_column_letter(col), get_column_letter(col + 3)
    ms.merge_cells(f"{c1}{fila}:{c2}{fila}")
    pon(ms, f"{c1}{fila}", etiqueta, negrita=True, tam=8, color=C_TENUE,
        fondo=C_TARJETA, alineacion="center")
    ms.merge_cells(f"{c1}{fila+1}:{c2}{fila+1}")
    pon(ms, f"{c1}{fila+1}", formula, negrita=True, tam=tam, color=color,
        fondo=C_TARJETA, fmt=fmt, alineacion="center")
    for f_ in (fila, fila + 1):
        for c_ in range(col, col + 4):
            ms.cell(row=f_, column=c_).border = CAJA
            ms.cell(row=f_, column=c_).fill = relleno(C_TARJETA)
    ms.row_dimensions[fila].height = 20
    ms.row_dimensions[fila + 1].height = 30
    if valor is not None:
        recordar(3, f"{c1}{fila+1}", valor)


tarjeta_kpi(5, 2, "💵  INGRESOS DEL MES", "=$AE$8", color=C_TEAL, valor=_ING)
tarjeta_kpi(5, 6, "🧾  GASTOS DEL MES", "=$AE$9", valor=_GAS)
tarjeta_kpi(5, 10, "⚖️  BALANCE", "=$AE$8-$AE$9", color=C_AZUL,
            valor=round(_ING - _GAS, 2))
tarjeta_kpi(5, 14, "📈  PROGRESO", "=$AE$14/100", fmt="0%", color=C_AZUL,
            valor=round(_PROGRESO / 100, 4))
tarjeta_kpi(5, 18, "🚦  ESTADO",
            '=IF($AE$13<=0,"🟢 Vas bien",IF($AE$14>=50,"🟡 Atención","🔴 En riesgo"))',
            fmt="General", tam=12,
            valor="🟢 Vas bien" if _FALTA <= 0 else ("🟡 Atención" if _PROGRESO >= 50
                                                    else "🔴 En riesgo"))

tarjeta_kpi(8, 2, "🎯  META DIARIA", "=$AE$15", color=C_VERDE, tam=19,
            valor=round(META, 2))
tarjeta_kpi(8, 6, "🗓️  DÍAS DE TRABAJO POR DELANTE", DIAS_DISPONIBLES, fmt="0")
tarjeta_kpi(8, 10, "📌  PAGOS SIN MARCAR", "=$AE$19", color=C_AMBAR,
            valor=_SIN_MARCAR)
tarjeta_kpi(8, 14, "🔁  GASTOS RECURRENTES DEL MES", "=$AE$10",
            valor=round(RECURRENTE_TOTAL, 2))
tarjeta_kpi(8, 18, "💳  TOTAL A CUBRIR", "=$AE$12", valor=_OBLIG)

ms.merge_cells("B11:V11")
if EN_RIESGO:
    aviso = (f"⚠️ EN RIESGO: {', '.join(ETIQUETAS_META)} vence el "
             f"{FECHA_META.strftime('%d/%m/%Y')} y no tienes ningún día de trabajo "
             f"marcado antes de esa fecha.")
    col_aviso, fondo_aviso = C_ROJO, C_ROJO_SUAVE
else:
    aviso = (f"🎯 Gana ${META:,.2f} cada día que trabajes para llegar a "
             f"{', '.join(ETIQUETAS_META)} el {FECHA_META.strftime('%d/%m/%Y')}.")
    col_aviso, fondo_aviso = C_TEXTO, C_TARJETA
pon(ms, "B11", aviso, negrita=True, tam=11, color=col_aviso, fondo=fondo_aviso,
    alineacion="center", ajustar=True, borde=True)
ms.row_dimensions[11].height = 34

# ---- Títulos de bloque ----
for rango, titulo in [("B14:H14", "A ·  REGISTRO DE MOVIMIENTOS"),
                      ("J14:N14", "B ·  PAGOS DEL MES"),
                      ("P14:V14", "C ·  CALENDARIO DE TRABAJO"),
                      ("X14:AB14", "D ·  PRESUPUESTO POR CATEGORÍA")]:
    ms.merge_cells(rango)
    pon(ms, rango.split(":")[0], titulo, negrita=True, tam=11, color=C_AZUL)
ms.row_dimensions[14].height = 26


def encabezados(fila, col_ini, textos):
    for i, t in enumerate(textos):
        c = ms.cell(row=fila, column=col_ini + i, value=t)
        c.font = Font(name=FUENTE, bold=True, size=9, color="FFFFFF")
        c.fill = relleno(C_TEXTO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CAJA
    ms.row_dimensions[fila].height = 32


# ---- A · Movimientos ----
encabezados(15, 2, ["Fecha", "Tipo", "Monto", "Descripción",
                    "Plataforma (sólo ingresos)", "Categoría (sólo gastos)"])
movimientos = []
for d, m, desc, plat in INGRESOS:
    movimientos.append((dt.date(ANIO, MES, d), "Ingreso", m, desc, plat, ""))
for d, m, desc, cat in GASTOS_MANUALES:
    movimientos.append((dt.date(ANIO, MES, d), "Gasto", m, desc, "", cat))
for nombre, monto, dia, tipo in LISTA_PAGOS:
    if nombre in PAGADOS:
        cat = "Pago de tarjeta de crédito" if tipo == "Tarjeta" else "Pago fijo"
        movimientos.append((dt.date(ANIO, MES, dia), "Gasto", monto, nombre, "", cat))
movimientos.sort(key=lambda m: (m[0], m[1]))

for i, fila_mov in enumerate(movimientos):
    for j, v in enumerate(fila_mov):
        ms.cell(row=FILA_MOV_1 + i, column=2 + j, value=v)
pintar(ms, f"B{FILA_MOV_1}:G{FILA_MOV_N}", fondo=C_INPUT, borde=True)
pintar(ms, f"B{FILA_MOV_1}:B{FILA_MOV_N}", fondo=C_INPUT, borde=True, fmt=FECHA)
pintar(ms, f"C{FILA_MOV_1}:C{FILA_MOV_N}", fondo=C_INPUT, borde=True, alineacion="center")
pintar(ms, f"D{FILA_MOV_1}:D{FILA_MOV_N}", fondo=C_INPUT, borde=True, fmt=DINERO)

# ---- B · Pagos del mes ----
encabezados(15, 10, ["¿Pagado?", "Concepto", "Monto", "Día", "Tipo"])
for i, (nombre, monto, dia, tipo) in enumerate(LISTA_PAGOS):
    f_ = FILA_PAG_1 + i
    ms.cell(row=f_, column=10, value="✔" if nombre in PAGADOS else "")
    ms.cell(row=f_, column=11, value=nombre)
    ms.cell(row=f_, column=12, value=monto)
    ms.cell(row=f_, column=13, value=dia)
    ms.cell(row=f_, column=14, value=tipo)
pintar(ms, f"J{FILA_PAG_1}:J{FILA_PAG_N}", fondo=C_INPUT, borde=True,
       alineacion="center", negrita=True, color=C_VERDE, tam=12)
pintar(ms, f"K{FILA_PAG_1}:N{FILA_PAG_N}", fondo="F5F8FC", borde=True)
pintar(ms, f"L{FILA_PAG_1}:L{FILA_PAG_N}", fondo="F5F8FC", borde=True, fmt=DINERO)
pintar(ms, f"M{FILA_PAG_1}:M{FILA_PAG_N}", fondo="F5F8FC", borde=True, alineacion="center")
pintar(ms, f"N{FILA_PAG_1}:N{FILA_PAG_N}", fondo="F5F8FC", borde=True,
       color=C_TENUE, tam=9, alineacion="center")
# Pagado en verde; vencido y sin marcar en rojo.
ms.conditional_formatting.add(
    f"K{FILA_PAG_1}:N{FILA_PAG_N}",
    FormulaRule(formula=[f'$J{FILA_PAG_1}="✔"'],
                fill=relleno(C_VERDE_SUAVE), font=Font(name=FUENTE, color=C_VERDE)))
ms.conditional_formatting.add(
    f"K{FILA_PAG_1}:N{FILA_PAG_N}",
    FormulaRule(formula=[f'AND($K{FILA_PAG_1}<>"",$J{FILA_PAG_1}="",'
                         f'$M{FILA_PAG_1}<{HOY.day})'],
                font=Font(name=FUENTE, bold=True, color=C_ROJO)))
ms.merge_cells(f"J{FILA_PAG_N+2}:N{FILA_PAG_N+2}")
pon(ms, f"J{FILA_PAG_N+2}",
    "En la plantilla real esto es una casilla: al marcarla, el gasto se registra solo.",
    color=C_TENUE, cursiva=True, ajustar=True)
ms.row_dimensions[FILA_PAG_N + 2].height = 30

# ---- C · Calendario ----
encabezados(15, 16, ["Dom", "Lun", "Mar", "Mié", "Jue", "Vie", "Sáb"])
primer_dia_semana = (dt.date(ANIO, MES, 1).weekday() + 1) % 7   # 0 = domingo
marcados = set(DIAS_TRABAJADOS + DIAS_PLANEADOS)
for semana in range(6):
    f_num = CAL_F1 + semana * 2
    f_chk = f_num + 1
    for col in range(7):
        indice = semana * 7 + col
        dia = indice - primer_dia_semana + 1
        cn = ms.cell(row=f_num, column=16 + col)
        cc = ms.cell(row=f_chk, column=16 + col)
        valido = 1 <= dia <= DIAS_MES
        cn.value = dia if valido else None
        cc.value = "✔" if (valido and dia in marcados) else None
        cn.border = cc.border = CAJA
        cn.alignment = cc.alignment = Alignment(horizontal="center", vertical="center")
        es_hoy = valido and dia == HOY.day
        cn.font = Font(name=FUENTE, bold=True, size=10,
                       color="FFFFFF" if es_hoy else C_TEXTO)
        cn.fill = relleno(C_AZUL if es_hoy else (C_TARJETA if valido else C_FONDO))
        cc.font = Font(name=FUENTE, bold=True, size=11, color=C_VERDE)
        cc.fill = relleno(C_INPUT if valido else C_FONDO)
    ms.row_dimensions[f_num].height = 18
    ms.row_dimensions[f_chk].height = 20
ms.merge_cells(f"P{CAL_F1+12}:V{CAL_F1+12}")
pon(ms, f"P{CAL_F1+12}",
    "Marca los días que VAS a trabajar. Los pasados se marcan solos al registrar un ingreso.",
    color=C_TENUE, cursiva=True, ajustar=True)
ms.row_dimensions[CAL_F1 + 12].height = 30

# ---- D · Presupuesto por categoría ----
encabezados(15, 24, ["Categoría", "Presupuesto", "Gastado", "%", "Avance"])
for i, (nombre, presupuesto) in enumerate(CATEGORIAS):
    f_ = PRE_F1 + i
    ms.cell(row=f_, column=24, value=nombre)
    ms.cell(row=f_, column=25, value=presupuesto)
    ms.cell(row=f_, column=26).value = (
        f'=SUMIFS($D${FILA_MOV_1}:$D${FILA_MOV_N},'
        f'$C${FILA_MOV_1}:$C${FILA_MOV_N},"Gasto",'
        f'$G${FILA_MOV_1}:$G${FILA_MOV_N},$X{f_})')
    ms.cell(row=f_, column=27).value = f'=IF($Y{f_}<=0,0,ROUND($Z{f_}/$Y{f_}*100,0))'
    ms.cell(row=f_, column=28).value = (
        f'=IF($X{f_}="","",REPT("█",MIN(20,ROUND($AA{f_}/5,0)))'
        f'&IF($AA{f_}>100," ⚠",""))')
    _gastado = round(sum(m for _, m, _, cat in GASTOS_MANUALES if cat == nombre), 2)
    _pct = 0 if presupuesto <= 0 else round(_gastado / presupuesto * 100)
    recordar(3, f"Z{f_}", _gastado)
    recordar(3, f"AA{f_}", _pct)
    recordar(3, f"AB{f_}", "█" * min(20, round(_pct / 5)) + (" ⚠" if _pct > 100 else ""))
pintar(ms, f"X{PRE_F1}:AB{PRE_FN}", fondo="F5F8FC", borde=True)
pintar(ms, f"Y{PRE_F1}:Z{PRE_FN}", fondo="F5F8FC", borde=True, fmt=DINERO)
pintar(ms, f"AA{PRE_F1}:AA{PRE_FN}", fondo="F5F8FC", borde=True, fmt='0"%"',
       alineacion="center")
for formula, color in [
        (f'AND($X{PRE_F1}<>"",$AA{PRE_F1}<80)', C_VERDE),
        (f'AND($X{PRE_F1}<>"",$AA{PRE_F1}>=80,$AA{PRE_F1}<100)', C_AMBAR),
        (f'AND($X{PRE_F1}<>"",$AA{PRE_F1}>=100)', C_ROJO)]:
    ms.conditional_formatting.add(
        f"AA{PRE_F1}:AB{PRE_FN}",
        FormulaRule(formula=[formula], font=Font(name=FUENTE, bold=True, color=color)))
ms.merge_cells(f"X{PRE_FN+2}:AB{PRE_FN+2}")
pon(ms, f"X{PRE_FN+2}",
    "Verde por debajo del 80 %, ámbar entre 80 % y 99 %, rojo al llegar al 100 %.",
    color=C_TENUE, cursiva=True, ajustar=True)
ms.row_dimensions[PRE_FN + 2].height = 30

# ---- Gráfico de dona ----
ms["AT1"], ms["AU1"] = "Categoría", "Total"
gasto_por_cat = {}
for _, _, m, _, _, cat in movimientos:
    if cat:
        gasto_por_cat[cat] = gasto_por_cat.get(cat, 0) + m
for i, (cat, total) in enumerate(sorted(gasto_por_cat.items(), key=lambda x: -x[1])):
    ms.cell(row=2 + i, column=46, value=cat)
    ms.cell(row=2 + i, column=47, value=round(total, 2))

dona = DoughnutChart(holeSize=55)
dona.title = "Gastos por categoría"
dona.height, dona.width = 7.5, 11.5
etiquetas = Reference(ms, min_col=46, min_row=2, max_row=1 + len(gasto_por_cat))
datos = Reference(ms, min_col=47, min_row=1, max_row=1 + len(gasto_por_cat))
dona.add_data(datos, titles_from_data=True)
dona.set_categories(etiquetas)
ms.add_chart(dona, f"P{CAL_F1+14}")

for col in range(29, 49):
    ms.column_dimensions[get_column_letter(col)].hidden = True
ms.column_dimensions["H"].hidden = True
ms.freeze_panes = "B16"

# ---- Listas desplegables ----
for rango, origen in [(f"C{FILA_MOV_1}:C{FILA_MOV_N}", '"Ingreso,Gasto"'),
                      (f"F{FILA_MOV_1}:F{FILA_MOV_N}", f'"{",".join(PLATAFORMAS)}"'),
                      (f"G{FILA_MOV_1}:G{FILA_MOV_N}",
                       f'"{",".join([c for c, _ in CATEGORIAS])},Pago fijo,'
                       f'Pago de tarjeta de crédito,Otros"')]:
    dv = DataValidation(type="list", formula1=origen, allow_blank=True, showDropDown=False)
    ms.add_data_validation(dv)
    dv.add(rango)

# ===================================================================== #
# Hoja: Plan de pago de deudas
# ===================================================================== #
ds = wb.create_sheet("Plan de pago de deudas")
lienzo(ds, 40, 10)
for col, a in {"B": 12, "C": 26, "D": 16, "E": 13, "F": 15, "G": 22}.items():
    ds.column_dimensions[col].width = a

ds.merge_cells("B2:G2")
pon(ds, "B2", "💳  Plan de pago de deudas", negrita=True, tam=18)
ds.row_dimensions[2].height = 40

pon(ds, "B4", "Estrategia", negrita=True)
pon(ds, "C4", "Bola de nieve", fondo=C_INPUT, borde=True)
pon(ds, "B5", "Pago extra mensual", negrita=True)
pon(ds, "C5", 100.00, fondo=C_INPUT, borde=True, fmt=DINERO)
dv = DataValidation(type="list", formula1='"Bola de nieve,Avalancha"', allow_blank=False,
                    showDropDown=False)
ds.add_data_validation(dv)
dv.add("C4")

ds.merge_cells("D4:G5")
pon(ds, "D4", "Bola de nieve = atacas primero la tarjeta con el saldo más chico.\n"
    "Avalancha = atacas primero la del interés más alto.",
    color=C_TENUE, ajustar=True)


def simular(tarjetas, estrategia, extra):
    """Misma simulación que apps-script/40_SimulacionDeuda.gs."""
    activas = [t for t in tarjetas if t[1] > 0]
    if not activas:
        return 0, 0.0, False, []
    orden = sorted(activas, key=(lambda t: t[1]) if estrategia == "bola"
                   else (lambda t: -t[2]))
    saldos = [t[1] for t in orden]
    liquidada = [None] * len(orden)
    presupuesto = sum(t[3] for t in orden) + max(0, extra)
    interes_total, meses, pendientes = 0.0, 0, len(orden)
    while pendientes > 0 and meses < 600:
        meses += 1
        for i, t in enumerate(orden):
            if saldos[i] <= 0:
                continue
            interes = saldos[i] * (t[2] / 100 / 12)
            interes_total += interes
            saldos[i] += interes
        objetivo = next((i for i in range(len(orden)) if saldos[i] > 0), None)
        if objetivo is None:
            break
        disponible = presupuesto
        for i, t in enumerate(orden):
            if i == objetivo or saldos[i] <= 0:
                continue
            pago = min(saldos[i], t[3])
            saldos[i] -= pago
            disponible -= pago
        saldos[objetivo] -= min(saldos[objetivo], max(0, disponible))
        for i in range(len(orden)):
            if saldos[i] <= 0.005 and liquidada[i] is None:
                saldos[i], liquidada[i] = 0, meses
                pendientes -= 1
    return meses, interes_total, pendientes > 0, [
        (t[0], t[1], t[2], t[3], liquidada[i] if liquidada[i] else meses)
        for i, t in enumerate(orden)]


meses, interes, nunca, orden_tarjetas = simular(TARJETAS, "bola", 100.0)

pon(ds, "B9", "Meses para liquidar todo", negrita=True)
pon(ds, "C9", meses, negrita=True, tam=14, color=C_AZUL, fondo="F5F8FC", borde=True, fmt="0")
pon(ds, "B10", "Interés total pagado", negrita=True)
pon(ds, "C10", round(interes, 2), negrita=True, tam=14, color=C_ROJO,
    fondo="F5F8FC", borde=True, fmt=DINERO)
pon(ds, "B11", "Fecha estimada de liquidación", negrita=True)
fecha_fin = dt.date(HOY.year + (HOY.month + meses - 1) // 12,
                    (HOY.month + meses - 1) % 12 + 1, 1)
pon(ds, "C11", fecha_fin, negrita=True, tam=14, fondo="F5F8FC", borde=True, fmt=FECHA)

anios, resto = divmod(meses, 12)
palabras = (f"{anios} año{'s' if anios != 1 else ''}" +
            (f" y {resto} meses" if resto else "")) if anios else f"{meses} meses"
ds.merge_cells("B13:G13")
pon(ds, "B13", f"✅ Con la estrategia «Bola de nieve» y $100.00 extra al mes, "
    f"liquidas {len(orden_tarjetas)} tarjetas en {palabras}.",
    negrita=True, color=C_VERDE, fondo=C_TARJETA, borde=True, ajustar=True,
    alineacion="center")
ds.row_dimensions[13].height = 34

for i, t in enumerate(["Orden", "Tarjeta", "Saldo actual", "APR %", "Pago mínimo",
                       "Meses hasta liquidarla"]):
    c = ds.cell(row=15, column=2 + i, value=t)
    c.font = Font(name=FUENTE, bold=True, size=9, color="FFFFFF")
    c.fill = relleno(C_TEXTO)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = CAJA
ds.row_dimensions[15].height = 32
for i, (nombre, saldo, apr, minimo, m) in enumerate(orden_tarjetas):
    f_ = 16 + i
    for j, v in enumerate([i + 1, nombre, saldo, apr, minimo, f"{m} meses"]):
        ds.cell(row=f_, column=2 + j, value=v)
pintar(ds, "B16:G25", fondo="F5F8FC", borde=True)
pintar(ds, "B16:B25", fondo="F5F8FC", borde=True, alineacion="center")
pintar(ds, "D16:D25", fondo="F5F8FC", borde=True, fmt=DINERO)
pintar(ds, "E16:E25", fondo="F5F8FC", borde=True, fmt='0.00"%"', alineacion="center")
pintar(ds, "F16:F25", fondo="F5F8FC", borde=True, fmt=DINERO)
pintar(ds, "G16:G25", fondo="F5F8FC", borde=True, alineacion="center")

ds.merge_cells("B27:G27")
pon(ds, "B27", "En la plantilla real un botón del menú recalcula esto con tus datos. "
    "Aquí los resultados son fijos, calculados con las tarjetas de ejemplo.",
    color=C_TENUE, cursiva=True, ajustar=True)
ds.row_dimensions[27].height = 30


def inyectar_cache(ruta, cache):
    """
    Añade el <v> calculado a cada celda de fórmula, sin tocar el <f>.

    openpyxl no escribe valores en caché, así que sin esto el visor de Google
    Drive muestra en blanco todo el panel hasta que el archivo se convierte a
    Google Sheets. Con esto se ve bien desde el primer vistazo, y la fórmula
    sigue ahí para recalcular.
    """
    import re
    import shutil
    import zipfile
    from xml.sax.saxutils import escape

    def procesar(xml, valores):
        def reemplazo(m):
            completo, ref = m.group(0), m.group(1)
            if ref not in valores or "<f" not in completo:
                return completo
            v = valores[ref]
            # openpyxl deja un <v/> vacío en las celdas de fórmula; hay que
            # quitarlo o gana él y la celda se lee como vacía.
            cuerpo = re.sub(r"<v\s*/>|<v>.*?</v>", "", completo, flags=re.S)
            if isinstance(v, str):
                if v == "":
                    return cuerpo
                cuerpo = re.sub(r'(<c\b[^>]*?)\s+t="[^"]*"', r"\1", cuerpo, count=1)
                cuerpo = cuerpo.replace("<c ", '<c t="str" ', 1)
                return cuerpo.replace("</c>", f"<v>{escape(str(v))}</v></c>")
            return cuerpo.replace("</c>", f"<v>{v}</v></c>")

        # Las celdas vacías se escriben autocerradas (<c r="D26" s="1"/>). Sin
        # contemplarlas, el patrón arranca en una de ellas y se traga las
        # celdas siguientes hasta el primer </c>, que ya es de otra celda.
        patron = r'<c r="([A-Z]+\d+)"[^>]*?(?:/>|>(?:(?!<c[ />]).)*?</c>)'
        return re.sub(patron, reemplazo, xml, flags=re.S)

    temporal = ruta + ".tmp"
    with zipfile.ZipFile(ruta) as zin, \
            zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", item.filename)
            if m and int(m.group(1)) in cache:
                datos = procesar(datos.decode("utf-8"), cache[int(m.group(1))]).encode("utf-8")
            zout.writestr(item, datos)
    shutil.move(temporal, ruta)


# --------------------------------------------------------------------- #
salida = str(pathlib.Path(__file__).with_name("Gig-Finance-vista-previa.xlsx"))
wb.save(salida)
inyectar_cache(salida, CACHE)
print(f"Guardado: {salida}")
print(f"Meta diaria calculada: ${META:,.2f} → {', '.join(ETIQUETAS_META)} "
      f"el {FECHA_META.strftime('%d/%m/%Y')} (en riesgo: {EN_RIESGO})")
print(f"Días disponibles: {DIAS_DISPONIBLES} · Recurrentes prorrateados: ${RECURRENTE_TOTAL:,.2f}")
print(f"Deuda: {meses} meses, interés ${interes:,.2f}")
